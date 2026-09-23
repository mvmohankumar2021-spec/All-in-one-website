from docx import Document
from pathlib import Path

doc=Document(); doc.add_heading('SHAKALPA Cleaning Safety Declaration',0)
doc.add_paragraph('This declaration confirms that the Cleaning Service Partner will provide safe, lawful, and professional house, deep, kitchen, bathroom, sofa, carpet, and water tank cleaning services on SHAKALPA.')
for h,t in [('Partner declaration','I confirm that my information and work proof are accurate and that I will undertake only work within my skills, experience, and applicable local permissions.'),('Customer and site safety','I will inspect the site before work begins, use suitable protective equipment and cleaning materials, keep the work area safe, and explain access, drying, ventilation, and hygiene requirements to the customer.'),('Cleaning work','I will use fit-for-purpose products, protect customer property, handle water and waste responsibly, and obtain customer approval before any material or scope change that affects price or timing.'),('Professional conduct','I will provide clear pricing, treat customers and their property respectfully, protect personal information, and comply with SHAKALPA policies and applicable laws.')]: doc.add_heading(h,1); doc.add_paragraph(t)
doc.add_paragraph('By accepting this declaration in the SHAKALPA partner application, I agree to follow these commitments.')
doc.save(Path('legal-documents')/'SHAKALPA_Cleaning_Safety_Declaration_India_Draft.docx')
