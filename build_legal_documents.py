from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent / "legal-documents"


def set_font(run, size=11, bold=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("SHAKALPA | Page ")
    set_font(run, 8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def paragraph(document, text="", bold_lead=None):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.18
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_font(lead, 11, True)
        rest = p.add_run(text[len(bold_lead):])
        set_font(rest)
    else:
        run = p.add_run(text)
        set_font(run)
    return p


def heading(document, text, level=1):
    p = document.add_paragraph()
    p.style = "Heading 1" if level == 1 else "Heading 2"
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run, 15 if level == 1 else 12, True)
    return p


def bullets(document, items):
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        run = p.add_run(item)
        set_font(run, 10.5)


def make_document(title, purpose, sections, acknowledgement, filename):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(11)
    for name in ("Heading 1", "Heading 2"):
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    header = section.header.paragraphs[0]
    header.text = "SHAKALPA"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_font(run, 9, True)
    add_page_number(section.footer.paragraphs[0])

    title_p = document.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_p.paragraph_format.space_after = Pt(5)
    title_run = title_p.add_run(title)
    set_font(title_run, 24, True)
    meta = document.add_paragraph()
    meta.paragraph_format.space_after = Pt(15)
    meta_run = meta.add_run("Version 1.0 | Partner onboarding draft | India")
    set_font(meta_run, 9)
    paragraph(document, purpose)
    paragraph(document, "This document is intended for partner onboarding and operational use. It should be reviewed and approved by SHAKALPA's legal and compliance teams before production use.")
    for section_title, intro, points in sections:
        heading(document, section_title)
        if intro:
            paragraph(document, intro)
        if points:
            bullets(document, points)
    heading(document, "Acknowledgement")
    paragraph(document, acknowledgement)
    paragraph(document, "By accepting this document in the SHAKALPA partner onboarding flow, the partner confirms that the information provided is accurate and agrees to comply with this document as updated from time to time.")
    OUT.mkdir(exist_ok=True)
    document.save(OUT / filename)


make_document(
    "SHAKALPA Code of Conduct",
    "This Code of Conduct sets the minimum standards for every SHAKALPA Service Partner, including electrical service partners, when serving customers through the platform.",
    [
        ("Scope and purpose", "Partners must conduct themselves professionally, safely, lawfully, and respectfully in every customer interaction.", []),
        ("Professional conduct", "Partners must:", ["provide accurate service descriptions, pricing, availability, qualifications, and contact information;", "arrive on time where an appointment is accepted, communicate delays promptly, and complete agreed work with reasonable care;", "treat customers, colleagues, and other partners with dignity and without discrimination, harassment, intimidation, or abusive language;", "not request payment outside SHAKALPA where platform rules require the platform payment flow."]),
        ("Customer safety and property", "Partners must:", ["use suitable tools, personal protective equipment, and safe work practices;", "obtain customer approval before material changes, additional work, or charges beyond the quoted scope;", "protect customer property, keep work areas orderly, and report damage or incidents promptly;", "stop work and escalate when a task cannot be performed safely or lawfully."]),
        ("Integrity and compliance", "Partners must not:", ["misrepresent licences, qualifications, customer ratings, photographs, service records, or prices;", "offer, request, or accept bribes, kickbacks, or improper benefits;", "use another person's identity, account, documents, or payment details without authority;", "perform work for which a licence, registration, insurance, or qualification is required unless the partner holds and maintains the applicable requirement."]),
        ("Privacy and reporting", "Partners must protect customer information and use it only to deliver the booked service. Suspected safety incidents, fraud, harassment, data misuse, or serious service failures must be reported to SHAKALPA without delay.", []),
        ("Enforcement", "SHAKALPA may investigate concerns, request evidence, pause listings or bookings, require corrective action, or suspend or terminate access where this Code or applicable law is breached.", []),
    ],
    "I have read and understood the SHAKALPA Code of Conduct and agree to follow it.",
    "SHAKALPA_Code_of_Conduct_India_Draft.docx",
)

make_document(
    "SHAKALPA Electrical Safety Declaration",
    "This Safety Declaration applies to Service Partners offering electrical repair, wiring, inspection, fan or light installation, inverter installation, or related electrical services through SHAKALPA.",
    [
        ("Purpose", "Electrical work can create risks to people and property. This declaration records the safety commitments required before a partner can be activated for electrical services.", []),
        ("Competence and authority", "The partner declares that they will:", ["offer only services within their competence, experience, and legal authority;", "maintain any licence, registration, certification, insurance, or supervision required for the specific work and location;", "not undertake higher-risk work such as home wiring, inverter installation, or electrical inspection without the appropriate qualification and registration details supplied to SHAKALPA."]),
        ("Safe work practices", "Before and during work, the partner will:", ["assess the work area and identify hazards before beginning;", "isolate, test, and verify electrical systems where required before working on them;", "use suitable insulated tools, testing equipment, protective equipment, and materials;", "keep customers, children, and bystanders away from hazards and explain any needed precautions;", "stop work when unsafe conditions, missing information, or required authorisations prevent safe completion."]),
        ("Pricing and materials", "The partner will explain inspection, labour, emergency, and material charges before incurring them where reasonably possible. Substitute materials or material costs outside the agreed scope require customer approval.", []),
        ("Incidents and escalation", "The partner will promptly notify emergency services when needed, make the area safe where possible, and report serious incidents, property damage, electrical hazards, or customer complaints to SHAKALPA.", []),
        ("No waiver of legal requirements", "This declaration does not replace any licence, statutory duty, local authority approval, electrical code, insurance requirement, or other legal obligation that applies to the work.", []),
    ],
    "I confirm that I will follow these safety commitments and will not accept or complete work that I am not qualified, authorised, or equipped to perform safely.",
    "SHAKALPA_Electrical_Safety_Declaration_India_Draft.docx",
)

make_document(
    "SHAKALPA Background Verification Consent",
    "This consent explains how SHAKALPA may conduct a background verification for a Service Partner before or after activation, where permitted by applicable law.",
    [
        ("What SHAKALPA may verify", "With the partner's consent and where lawful, SHAKALPA may verify identity information, address information, work experience or credentials, licences or registrations, business details, public professional references, and records relevant to safety, fraud prevention, or platform trust.", []),
        ("Purpose", "Verification supports customer safety, platform integrity, prevention of fraud, and compliance with SHAKALPA's onboarding standards. It is not a guarantee of employment, bookings, earnings, or continued activation.", []),
        ("Sources and sharing", "SHAKALPA may use information supplied by the partner, authorised verification providers, issuing authorities, professional references, and publicly available sources where permitted. Information is shared only with personnel and providers who need it for verification, compliance, or safety purposes.", []),
        ("Partner responsibilities", "The partner must provide accurate information, promptly disclose material changes to verification information, and cooperate with reasonable document or clarification requests. False, incomplete, or misleading information may delay, suspend, or end activation.", []),
        ("Privacy and retention", "SHAKALPA will handle verification information under its Privacy Policy and retain it only for as long as needed for the stated purposes, legal obligations, dispute handling, or legitimate security needs.", []),
        ("Review and questions", "Where appropriate, a partner may ask SHAKALPA to correct inaccurate information or request an explanation of a verification outcome, subject to legal, safety, and fraud-prevention limits.", []),
    ],
    "I voluntarily consent to the background verification activities described in this document and confirm that I have authority to provide the information and documents requested.",
    "SHAKALPA_Background_Verification_Consent_India_Draft.docx",
)

make_document(
    "SHAKALPA Home Technical Services Safety Declaration",
    "This Safety Declaration applies to Service Partners offering locksmith, waterproofing, roof repair, CCTV installation, solar installation, or packers and movers services through SHAKALPA.",
    [
        ("Competence and legal requirements", "Partners must offer only work they are competent, equipped, authorised, insured, and legally permitted to perform. Any required licence, registration, certification, local approval, or supervision must be kept current.", []),
        ("Safe work and customer property", "Before starting work, partners must assess risks, use suitable tools and protective equipment, protect customer property, explain material risks, and stop work where it cannot be completed safely.", ["For roof, waterproofing, CCTV, and solar work, use safe access methods and do not work in unsafe weather or hazardous electrical conditions.", "For locksmith work, verify the customer's authority to request access before opening, changing, or bypassing a lock.", "For packers and movers work, use safe lifting, packing, loading, transport, and delivery procedures, and record any pre-existing or newly observed damage promptly."]),
        ("Pricing, warranties, and claims", "Partners must give clear information about inspection, labour, materials, warranty or revisit terms, and any damage-claim process before undertaking chargeable work. Additional charges or scope changes require customer approval.", []),
        ("Incidents and escalation", "Partners must report serious safety incidents, damage, loss, security concerns, or customer complaints to SHAKALPA without delay and cooperate with reasonable follow-up requests.", []),
    ],
    "I confirm that I will meet these safety commitments and will not accept work that I cannot perform safely, lawfully, and professionally.",
    "SHAKALPA_Home_Technical_Safety_Declaration_India_Draft.docx",
)

make_document(
    "SHAKALPA Construction Safety Declaration",
    "This Safety Declaration applies to Service Partners offering residential or commercial construction, renovation, civil contracting, structural work, or site supervision through SHAKALPA.",
    [
        ("Competence and authority", "The partner must undertake only work that it is competent, licensed, insured, staffed, and legally authorised to perform. Structural work and site supervision must be performed or supervised by appropriately qualified professionals where required.", []),
        ("Site safety", "The partner must assess site risks before work begins, use suitable PPE and equipment, maintain safe access and housekeeping, brief workers on hazards, and stop work when unsafe conditions arise.", ["Follow applicable building, labour, fire, electrical, environmental, and local-authority requirements.", "Protect customers, occupants, visitors, neighbouring property, and workers from foreseeable risks.", "Report serious incidents, injuries, unsafe conditions, material damage, or suspected structural risks to SHAKALPA without delay."]),
        ("Scope, pricing, and changes", "The partner must provide a clear scope, quotation, materials policy, timeline, and change-order process. Any material change to scope, price, or design requires customer approval before work proceeds.", []),
        ("Quality, warranty, and claims", "The partner must state the applicable warranty or defect-liability period and maintain a fair documented process for damage claims, remedial work, and dispute handling.", []),
    ],
    "I confirm that I will comply with these construction safety commitments and will not accept work that cannot be completed safely, lawfully, and professionally.",
    "SHAKALPA_Construction_Safety_Declaration_India_Draft.docx",
)

make_document(
    "SHAKALPA Architectural & Engineering Compliance Declaration",
    "This declaration applies to Service Partners providing architectural design, civil engineering, structural engineering, building-plan approval support, or 3D elevation design through SHAKALPA.",
    [
        ("Professional competence and registration", "Partners must provide only services within their competence, professional registration, licence, insurance, and legal authority. Where a law, code, or authority requires a registered architect or engineer, the partner must ensure that the work is prepared, checked, or supervised by an appropriately qualified professional.", []),
        ("Design quality and safety", "Partners must use reasonable professional care, accurately communicate assumptions and limitations, identify material design or site risks, and recommend that work stop or be reviewed when safety, structural integrity, approvals, or required information are uncertain.", []),
        ("Plans, approvals, and privacy", "Partners must explain the scope of drawings, approval support, authority submissions, expected timelines, revision limits, and third-party dependencies. Customer plans, site details, and documents must be used only for providing the agreed service and handled in accordance with the Privacy Policy.", []),
        ("Pricing, revisions, and disputes", "Partners must provide clear consultation, site-visit, drawing-package, revision, and additional-work pricing. Changes to scope or fees require customer approval. Partners must maintain a fair process for correction requests, disputes, and professional complaints.", []),
    ],
    "I confirm that I will comply with these professional and compliance commitments and will not accept work that I am not qualified, registered, insured, or authorised to perform.",
    "SHAKALPA_Architectural_Engineering_Compliance_Declaration_India_Draft.docx",
)

make_document(
    "SHAKALPA Interior Design Safety & Compliance Declaration",
    "This declaration applies to Service Partners providing home or office interior design, modular kitchens, wardrobe design, false ceilings, or space planning through SHAKALPA.",
    [
        ("Professional scope and competence", "Partners must undertake only work they are qualified, insured, equipped, and authorised to provide. They must accurately explain whether they provide design only, material supply, installation, project coordination, or execution supervision.", []),
        ("Customer property, site safety, and materials", "Partners must use reasonable care when visiting customer premises, coordinating contractors, selecting materials, planning false ceilings, furniture, wardrobes, kitchens, or fixtures. Safety risks, material limitations, and access requirements must be explained before work begins.", []),
        ("Quotations, changes, and timelines", "Partners must provide clear deliverables, material policy, estimated timeline, quotation, revision limits, and change-order process. Material substitutions, additional charges, or scope changes require customer approval.", []),
        ("Warranty, damage claims, and privacy", "Partners must state applicable warranty or defect-liability terms and maintain a fair documented process for damage claims, remedial work, and disputes. Customer drawings, measurements, property images, and contact information must be handled under the Privacy Policy.", []),
    ],
    "I confirm that I will follow these interior design safety and compliance commitments and will not accept work that cannot be performed safely, lawfully, and professionally.",
    "SHAKALPA_Interior_Design_Compliance_Declaration_India_Draft.docx",
)

make_document(
    "SHAKALPA Flooring & Wall-Cladding Safety Declaration",
    "This declaration applies to Service Partners providing tile installation, marble, granite, wooden, or epoxy flooring, or wall cladding through SHAKALPA.",
    [
        ("Trade competence and equipment", "Partners must undertake only flooring and cladding work that their team is trained, equipped, insured, and authorised to perform. Specialist stone, epoxy, and cladding work requires suitable work proof and safe cutting, handling, polishing, and installation equipment.", []),
        ("Site safety and property protection", "Partners must assess the surface and work area before starting, use appropriate PPE and dust controls, manage cutting and chemical hazards, protect customer property, and report unsafe conditions, injuries, or damage promptly.", []),
        ("Materials, pricing, and change orders", "Partners must clearly explain the material policy, surface preparation, levelling, waterproofing, grouting, polishing, sealing, skirting, removal, disposal, transport, and any additional charges. Material substitutions and changes require customer approval.", []),
        ("Quality, warranty, and claims", "Partners must state applicable workmanship warranty or defect-liability terms and maintain a fair process for revisit work, breakage, wastage, damage claims, and dispute resolution.", []),
    ],
    "I confirm that I will comply with these flooring and wall-cladding safety commitments and will not accept work that cannot be completed safely, lawfully, and professionally.",
    "SHAKALPA_Flooring_Cladding_Safety_Declaration_India_Draft.docx",
)

make_document(
    "SHAKALPA Fabrication & Metalwork Safety Declaration",
    "This declaration applies to Service Partners providing steel or aluminium fabrication, glass work, welding, gate fabrication, or grill work through SHAKALPA.",
    [
        ("Trade competence and workshop safety", "Partners must undertake only work their team is trained, equipped, insured, and authorised to perform. Welding, cutting, grinding, lifting, glass handling, and installation must use suitable equipment, PPE, electrical controls, and fire-safety measures.", []),
        ("Site safety and customer property", "Partners must inspect the work area, protect people and property, control sparks, dust, sharp edges, lifting risks, and glass hazards, and stop work when conditions are unsafe. Incidents, injuries, or material damage must be reported promptly.", []),
        ("Scope, materials, and pricing", "Partners must clearly state measurements, design approval, material type and grade, finishing, coating, glass fitting, transport, installation, exclusions, quotation, and change-order process. Changes require customer approval.", []),
        ("Warranty and claims", "Partners must state workmanship warranty, revisit terms, and applicable exclusions for corrosion, rust, coating, or glass breakage, and maintain a fair process for damage claims and disputes.", []),
    ],
    "I confirm that I will comply with these fabrication and metalwork safety commitments and will not accept work that cannot be performed safely, lawfully, and professionally.",
    "SHAKALPA_Fabrication_Metalwork_Safety_Declaration_India_Draft.docx",
)

make_document(
    "SHAKALPA Building Materials Supply Quality & Safety Declaration",
    "This declaration applies to Service Partners supplying cement, steel, bricks and blocks, sand and aggregates, tiles and sanitaryware, hardware, or doors and windows through SHAKALPA.",
    [
        ("Quality, authorisation, and traceability", "Suppliers must provide genuine, accurately described products from lawful sources and maintain appropriate supplier, dealer, distributor, manufacturer, GST, licence, brand, grade, test-certificate, and source documentation where applicable.", []),
        ("Stock, delivery, and handling", "Suppliers must give accurate stock, lead-time, minimum-order, transport, loading, unloading, and delivery information. Materials must be handled and delivered with reasonable care to reduce foreseeable damage, shortage, contamination, or safety risks.", []),
        ("Pricing, substitutions, returns, and claims", "Suppliers must clearly communicate unit pricing, taxes, transport charges, material substitutions, returns, replacements, refunds, warranty support, and procedures for shortage, breakage, damage claims, and delivery disputes. Substitutions require customer approval.", []),
        ("Invoices and professional conduct", "Suppliers must issue accurate invoices and use customer information only to fulfil the agreed order. Suspected fraud, unsafe materials, serious delivery damage, or customer complaints must be reported promptly to SHAKALPA.", []),
    ],
    "I confirm that I will comply with these building-materials quality and safety commitments and will supply products professionally, lawfully, and with accurate supporting information.",
    "SHAKALPA_Building_Materials_Supply_Quality_Safety_Declaration_India_Draft.docx",
)

make_document(
    "SHAKALPA Real Estate Sales Compliance Declaration",
    "This declaration applies to Service Partners offering residential or commercial property sales, land and plot sales, new projects, or resale property services through SHAKALPA.",
    [
        ("Registration, authority, and accurate listings", "Partners must maintain required RERA, brokerage, GST, business, insurance, builder/developer, owner/seller, and listing authorisations. Listings must accurately describe property ownership status, approvals, area, price, project details, and material limitations.", []),
        ("Due diligence and customer safety", "Partners must follow a reasonable process for title, ownership, encumbrance, RERA/project approval, and listing-document review. They must use safe property-visit practices, protect customer information, and not make misleading claims or collect unauthorised payments.", []),
        ("Commissions, advertising, and consent", "Partners must clearly disclose consultation, site-visit, brokerage, commission, cancellation, refund, advertising, and lead-handling terms. Property photos, listing data, and customer contact details must be used with appropriate authority and in accordance with the Privacy Policy.", []),
        ("Complaints and escalation", "Partners must maintain a fair grievance and dispute process and promptly report suspected fraud, forged documents, serious safety concerns, or material listing inaccuracies to SHAKALPA.", []),
    ],
    "I confirm that I will comply with these real-estate sales compliance commitments and will provide only authorised, accurate, and professionally handled property services.",
    "SHAKALPA_Real_Estate_Sales_Compliance_Declaration_India_Draft.docx",
)

make_document(
    "SHAKALPA Real Estate Rental Compliance Declaration",
    "This declaration applies to Service Partners offering house, apartment, commercial, office, shop, or warehouse rental services through SHAKALPA.",
    [
        ("Registration and listing authority", "Partners must maintain applicable brokerage, RERA, GST, business, insurance, owner, property-management, and rental-listing authorisations. Rental listings must accurately state property type, rent, deposit, lease term, availability, use restrictions, and material conditions.", []),
        ("Tenant screening and viewing safety", "Partners must use a reasonable tenant KYC, screening, identity-verification, viewing, property-access, and customer-safety process. Customer and property information must be handled in accordance with the Privacy Policy.", []),
        ("Lease, deposits, maintenance, and disputes", "Partners must clearly disclose brokerage, cancellation, refund, rental-agreement, renewal, notice, deposit, inventory, handover, maintenance, repair, emergency-contact, damage-claim, and grievance processes. Commercial rentals require appropriate use approvals and lease documentation.", []),
        ("Professional conduct", "Partners must not make misleading claims, collect unauthorised payments, or share property or customer information without authority. Suspected fraud, forged documents, unsafe visits, or serious complaints must be reported promptly to SHAKALPA.", []),
    ],
    "I confirm that I will comply with these real-estate rental compliance commitments and will provide only authorised, accurate, and professionally managed rental services.",
    "SHAKALPA_Real_Estate_Rental_Compliance_Declaration_India_Draft.docx",
)

make_document(
    "SHAKALPA Accommodation Safety & Guest Compliance Declaration",
    "This declaration applies to Service Partners offering PG accommodation, hostels, service apartments, or guest houses through SHAKALPA.",
    [
        ("Authority and accommodation compliance", "Partners must own, operate, or manage listed properties under valid authority and maintain applicable lodging, local authority, business, GST, lease, management, and safety documentation. Listings must accurately describe rooms, beds, occupancy, restrictions, pricing, amenities, and availability.", []),
        ("Guest safety, security, and hygiene", "Partners must maintain reasonable fire safety, emergency exits, CCTV or security controls where applicable, first aid, hygiene, housekeeping, maintenance, and emergency-contact processes. They must use a reasonable guest KYC, check-in, visitor, access, and incident-escalation process.", []),
        ("Pricing, deposits, rules, and claims", "Partners must clearly disclose tariffs, rent, deposits, maintenance, food, utilities, minimum stay, cancellation, refund, house rules, damage, claims, eviction, termination, and dispute processes before accepting a booking or tenancy.", []),
        ("Privacy and conduct", "Guest and property information must be handled in accordance with the Privacy Policy. Serious safety incidents, suspected fraud, harassment, or significant property damage must be reported promptly to SHAKALPA.", []),
    ],
    "I confirm that I will comply with these accommodation safety and guest-compliance commitments and will provide only authorised, accurate, safe, and professionally managed accommodation services.",
    "SHAKALPA_Accommodation_Safety_Guest_Compliance_Declaration_India_Draft.docx",
)

make_document(
    "SHAKALPA Real Estate Services Compliance Declaration",
    "This declaration applies to Service Partners offering real-estate agency, property management, property valuation, property legal support, or home-loan assistance through SHAKALPA.",
    [
        ("Registration, authority, and professional scope", "Partners must maintain applicable RERA, brokerage, business, GST, professional, valuation, legal-practice, lender/DSA, insurance, owner, and property-management authorisations. Partners must accurately state their role and must not provide regulated legal, valuation, or lending advice unless properly authorised.", []),
        ("Property information and customer protection", "Partners must use reasonable due diligence for property, owner, title, approval, valuation, lender, and listing information within their professional scope. Fees, commissions, lender relationships, conflicts, limitations, cancellation, refund, damage-claim, and grievance processes must be disclosed clearly before engagement.", []),
        ("Management, legal, valuation, and loan assistance", "Property-management partners must maintain documented inspection, tenant or occupant, rent, maintenance, vendor, repair, emergency, and owner-reporting processes. Valuation reports must state their purpose and method. Legal services must be delivered only by suitably qualified professionals. Loan assistance must not promise approval or collect unauthorised charges.", []),
        ("Privacy, conduct, and escalation", "Customer, owner, tenant, property, financial, and identity information must be handled in accordance with the Privacy Policy. Partners must not make misleading claims, conceal material conflicts, or collect unauthorised payments, and must promptly report suspected fraud, forged documents, major safety concerns, or serious complaints to SHAKALPA.", []),
    ],
    "I confirm that I will comply with these real-estate services commitments and will provide only authorised, accurate, transparent, and professionally handled services.",
    "SHAKALPA_Real_Estate_Services_Compliance_Declaration_India_Draft.docx",
)

make_document(
    "SHAKALPA Marketplace Payments, Cancellation & Refund Policy",
    "This policy explains the marketplace payment, fulfilment, cancellation, and refund expectations for Service Partners using SHAKALPA payment collection and vendor settlement.",
    [
        ("Payment collection and settlement", "SHAKALPA or its payment provider may collect customer payments and verify payout details before settlement. Partners must provide accurate payout and tax information and must not request unauthorised off-platform payments for marketplace orders or bookings.", []),
        ("Fulfilment and service delivery", "Partners must accurately state their fulfilment method, availability, scope, charges, and delivery or service timelines. Partners must promptly notify customers and SHAKALPA of material delays, cancellations, unavailable items, or changes that affect an order or booking.", []),
        ("Cancellation and refunds", "Partners must publish clear cancellation and refund terms before a customer pays. Refund eligibility, deductions, timing, and any rescheduling or replacement process must be fair, accurately described, and handled in accordance with applicable law, payment-provider rules, and the customer-facing policy.", []),
        ("Disputes, chargebacks, and records", "Partners must reasonably cooperate with payment disputes, chargebacks, refund requests, and fraud reviews. They must retain relevant order, fulfilment, communication, and refund records and must not misrepresent completion or delivery.", []),
    ],
    "I confirm that I have reviewed this Marketplace Payments, Cancellation & Refund Policy and will follow it when using SHAKALPA marketplace payment collection and settlement.",
    "SHAKALPA_Marketplace_Payments_Cancellation_Refund_Policy_India_Draft.docx",
)

make_document(
    "SHAKALPA Bakery Food Safety & Quality Declaration",
    "This declaration applies to Service Partners offering bakery, cake shop, sweet shop, or dessert shop services through SHAKALPA.",
    [
        ("Food business compliance", "Partners must maintain applicable FSSAI, business, GST, facility, and local registration records and accurately describe products, ingredients, prices, and availability.", []),
        ("Food safety and allergens", "Partners must maintain safe preparation, storage, hygiene, packaging, temperature, labelling, expiry, allergen, cross-contamination, pest-control, and waste-handling practices.", []),
        ("Orders, delivery, and customer protection", "Partners must clearly disclose order lead time, customisation, delivery, cancellation, refund, replacement, damage-claim, and complaint processes and promptly escalate material food-quality or safety incidents.", []),
    ],
    "I confirm that I will provide safe, accurately described, and professionally handled bakery and sweets services.",
    "SHAKALPA_Bakery_Food_Safety_Quality_Declaration_India_Draft.docx",
)

make_document(
    "SHAKALPA Restaurant Food Safety & Quality Declaration",
    "This declaration applies to restaurant partners offering vegetarian, non-vegetarian, multi-cuisine, regional, Chinese, or fast-food services through SHAKALPA.",
    [
        ("Food business compliance", "Partners must maintain applicable FSSAI, business, GST, facility, and local registration records and accurately describe menu items, ingredients, prices, availability, and dietary claims.", []),
        ("Food safety and allergens", "Partners must maintain safe preparation, storage, hygiene, temperature, labelling, allergen, vegetarian/non-vegetarian separation, pest-control, and waste-handling practices.", []),
        ("Orders and customer protection", "Partners must clearly disclose fulfilment, delivery, cancellation, refund, replacement, complaint, and food-quality incident processes and promptly escalate material safety concerns.", []),
    ],
    "I confirm that I will provide safe, accurately described, and professionally handled restaurant services.",
    "SHAKALPA_Restaurant_Food_Safety_Quality_Declaration_India_Draft.docx",
)

make_document(
    "SHAKALPA Café & Beverages Food Safety & Quality Declaration",
    "This declaration applies to café, tea shop, juice shop, ice cream shop, snack shop, and street-food partners through SHAKALPA.",
    [("Food safety and quality", "Partners must maintain applicable FSSAI and business records; safe water, ice, ingredient, allergen, storage, refrigeration, hygiene, packaging, delivery, pest-control, and waste practices; and clear customer refund and complaint processes.", [])],
    "I confirm that I will provide safe, accurately described, and professionally handled café and beverages services.",
    "SHAKALPA_Cafe_Beverages_Food_Safety_Quality_Declaration_India_Draft.docx",
)
