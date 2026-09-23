from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from pathlib import Path

output = Path('legal-documents') / 'SHAKALPA_Furniture_Safety_Declaration_India_Draft.docx'
doc = Document()
section = doc.sections[0]
section.top_margin = Pt(54)
section.bottom_margin = Pt(54)
title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('SHAKALPA Furniture Safety Declaration')
intro = doc.add_paragraph('This declaration confirms that the Furniture Service Partner will provide safe, lawful, and professional furniture, door, window, wardrobe, and modular installation services on SHAKALPA.')
intro.paragraph_format.space_after = Pt(12)
for heading, body in [
    ('Partner declaration', 'I confirm that the information and work proof I provide are accurate. I will undertake only work within my skills, experience, and applicable local permissions.'),
    ('Customer and site safety', 'I will assess the site before work begins, use suitable tools and protective equipment, keep the work area safe, and explain material, repair, and installation limitations to the customer.'),
    ('Materials and installation', 'I will use fit-for-purpose materials, follow manufacturer guidance where applicable, and obtain customer approval before any material or scope change that affects price or timing.'),
    ('Professional conduct', 'I will treat customers and their property respectfully, protect personal information, provide clear pricing, and comply with SHAKALPA policies and applicable laws.'),
]:
    doc.add_heading(heading, level=1)
    doc.add_paragraph(body)
doc.add_paragraph('By accepting this declaration in the SHAKALPA partner application, I agree to follow these commitments.')
doc.save(output)
