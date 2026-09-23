from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from pathlib import Path

output = Path('legal-documents') / 'SHAKALPA_Painting_Safety_Declaration_India_Draft.docx'
doc = Document(); section = doc.sections[0]; section.top_margin = Pt(54); section.bottom_margin = Pt(54)
title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER; title.add_run('SHAKALPA Painting Safety Declaration')
doc.add_paragraph('This declaration confirms that the Painting Service Partner will provide safe, lawful, and professional interior, exterior, texture, waterproof, wood, and metal painting services on SHAKALPA.')
for heading, text in [('Partner declaration','I confirm that my information and work proof are accurate. I will undertake only work within my skills, experience, and applicable local permissions.'),('Customer and site safety','I will inspect the site before work begins, use suitable protective equipment, keep the work area safe, and communicate ventilation, access, height, and drying-time risks to the customer.'),('Paint and surface work','I will use fit-for-purpose materials, follow manufacturer guidance, protect nearby property, and obtain customer approval before any material or scope change that affects price or timing.'),('Professional conduct','I will treat customers and their property respectfully, provide clear pricing, protect personal information, and comply with SHAKALPA policies and applicable laws.')]:
    doc.add_heading(heading, level=1); doc.add_paragraph(text)
doc.add_paragraph('By accepting this declaration in the SHAKALPA partner application, I agree to follow these commitments.')
doc.save(output)
